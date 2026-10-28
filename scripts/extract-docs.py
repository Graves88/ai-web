# 从原始 PPTX / XLSX 抽取文本，按"页/幻灯片"切分，供后续 RAG 入库
# 输出：data/extracted/<文件名>.json
#   { source, type, pages: [{ page, slide_title, text }], totalChars, emptyPages }
# 说明：产品资料多为图片型 PPT，纯图页会 text 为空（后续可选 OCR）
import glob
import json
import os
import re

ROOT = r"D:\deepseek work\data"
SRC = os.path.join(ROOT, "待导入", "原始文件")
OUT = os.path.join(ROOT, "extracted")
os.makedirs(OUT, exist_ok=True)


def clean(text: str) -> str:
    text = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{2,}", "\n", text)
    return text.strip()


def extract_pptx(path):
    from pptx import Presentation

    prs = Presentation(path)
    pages = []
    for idx, slide in enumerate(prs.slides, start=1):
        chunks = []
        for shape in slide.shapes:
            try:
                if shape.has_text_frame:
                    t = shape.text_frame.text
                    if t and t.strip():
                        chunks.append(t)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                        if cells:
                            chunks.append(" | ".join(cells))
            except Exception:
                continue
        text = clean("\n".join(chunks))
        first_line = text.split("\n")[0][:40] if text else ""
        pages.append({"page": idx, "slide_title": first_line, "text": text})
    return pages


def extract_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    pages = []
    for idx, ws in enumerate(wb.worksheets, start=1):
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        text = clean("\n".join(rows))
        pages.append({"page": idx, "slide_title": ws.title, "text": text})
    return pages


total_docs = 0
summary = []
for path in sorted(glob.glob(os.path.join(SRC, "*"))):
    name = os.path.basename(path)
    ext = os.path.splitext(name)[1].lower()
    try:
        if ext == ".pptx":
            pages = extract_pptx(path)
            kind = "pptx"
        elif ext == ".xlsx":
            pages = extract_xlsx(path)
            kind = "xlsx"
        else:
            continue
    except Exception as e:
        print(f"跳过 {name}: {e}")
        continue

    total_chars = sum(len(p["text"]) for p in pages)
    empty = sum(1 for p in pages if not p["text"])
    out = {
        "source": name,
        "type": kind,
        "pages": pages,
        "totalChars": total_chars,
        "emptyPages": empty,
    }
    out_path = os.path.join(OUT, os.path.splitext(name)[0] + ".json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    total_docs += 1
    summary.append((name, len(pages), total_chars, empty))
    print(f"{name}: {len(pages)} 页, {total_chars} 字, 空页 {empty}")

print(f"\n共抽取 {total_docs} 份文档 -> {OUT}")
print("纯图页比例较高说明需要 OCR；否则可直接切分入库。")
