# 从在售清单提取"在售状态"汇总（非"在售"的型号），作为知识库的一个补充文档
# 输出 data/extracted/产品在售状态汇总（自动生成）.json，供 build-kb 入库
import json
import os
from collections import defaultdict
from openpyxl import load_workbook

SRC = r"D:\deepseek work\data\待导入\原始文件\在售清单_欧普家用智能照明在售清单-（共享清单）-26年10月 (1).xlsx"
OUT_DIR = r"D:\deepseek work\data\extracted"
os.makedirs(OUT_DIR, exist_ok=True)

wb = load_workbook(SRC, data_only=True, read_only=True)
groups = defaultdict(list)
total_rows = 0

for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    hdr_idx = None
    for i, r in enumerate(rows[:6]):
        if r and any(x is not None and str(x).strip() == "是否在售" for x in r):
            hdr_idx = i
            break
    if hdr_idx is None:
        continue
    header = [str(x).strip() if x is not None else "" for x in rows[hdr_idx]]
    status_col = header.index("是否在售")
    col = {name: i for i, name in enumerate(header) if name}

    for r in rows[hdr_idx + 1 :]:
        if not r or len(r) <= status_col or r[status_col] is None:
            continue
        status = str(r[status_col]).strip()
        if not status or status == "是否在售":
            continue
        total_rows += 1
        if status == "在售":
            continue

        def get(name):
            i = col.get(name)
            if i is None or i >= len(r) or r[i] is None:
                return ""
            return str(r[i]).strip()

        series = get("产品系列")
        code = get("物料代码")
        desc = get("产品描述")
        line = f"- [{ws.title}] {series}｜{code}｜{desc}".rstrip("｜")
        if line not in groups[status]:
            groups[status].append(line)

lines = ["在售清单中非「在售」状态的型号汇总（用于回答“哪些型号预退/待上市/何时上市”）", ""]
for status, items in sorted(groups.items(), key=lambda kv: -len(kv[1])):
    lines.append(f"【{status}】共 {len(items)} 个型号：")
    lines.extend(items[:30])
    lines.append("")

text = "\n".join(lines).strip()
doc = {
    "source": "产品在售状态汇总（自动生成）.json",
    "type": "xlsx",
    "pages": [{"page": 1, "slide_title": "在售状态汇总", "text": text}],
    "totalChars": len(text),
    "emptyPages": 0,
}
out = os.path.join(OUT_DIR, "产品在售状态汇总（自动生成）.json")
with open(out, "w", encoding="utf-8") as f:
    json.dump(doc, f, ensure_ascii=False, indent=1)

print("数据行:", total_rows)
for s, items in groups.items():
    print(f"  {s}: {len(items)} 个")
print("输出:", out)
